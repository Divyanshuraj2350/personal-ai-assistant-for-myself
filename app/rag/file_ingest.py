from pathlib import Path

from app.rag.ingest import ingest_text


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
}


# ============================================================
# READ TXT
# ============================================================

def extract_text_from_txt(file_path):
    """
    Extract text from a TXT file.
    """

    path = Path(file_path)

    if not path.exists():
        return {
            "status": "error",
            "error": "File does not exist.",
        }

    if not path.is_file():
        return {
            "status": "error",
            "error": "Path is not a file.",
        }

    try:
        text = path.read_text(
            encoding="utf-8"
        )

    except UnicodeDecodeError:
        return {
            "status": "error",
            "error": "The TXT file is not UTF-8 encoded.",
        }

    except OSError as error:
        return {
            "status": "error",
            "error": str(error),
        }

    if not text.strip():
        return {
            "status": "error",
            "error": "The TXT file is empty.",
        }

    return {
        "status": "success",
        "text": text,
    }


# ============================================================
# READ PDF
# ============================================================

def extract_text_from_pdf(file_path):
    """
    Extract text from a PDF file.

    Page markers are preserved so the AI can
    understand where information came from.
    """

    path = Path(file_path)

    if not path.exists():
        return {
            "status": "error",
            "error": "File does not exist.",
        }

    if not path.is_file():
        return {
            "status": "error",
            "error": "Path is not a file.",
        }

    try:
        from pypdf import PdfReader

        reader = PdfReader(
            str(path)
        )

        if not reader.pages:
            return {
                "status": "error",
                "error": "The PDF contains no pages.",
            }

        pages = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            if text.strip():
                pages.append(
                    f"\n[Page {page_number}]\n{text}"
                )

        full_text = "\n".join(
            pages
        ).strip()

    except Exception as error:
        return {
            "status": "error",
            "error": (
                f"Could not read PDF: {error}"
            ),
        }

    if not full_text:
        return {
            "status": "error",
            "error": (
                "No extractable text was found "
                "in the PDF. It may be a "
                "scanned/image-only PDF."
            ),
        }

    return {
        "status": "success",
        "text": full_text,
    }


# ============================================================
# READ DOCX
# ============================================================

def extract_text_from_docx(file_path):
    """
    Extract text from a DOCX file.

    Paragraphs and table contents are included.
    """

    path = Path(file_path)

    if not path.exists():
        return {
            "status": "error",
            "error": "File does not exist.",
        }

    if not path.is_file():
        return {
            "status": "error",
            "error": "Path is not a file.",
        }

    try:
        from docx import Document

        document = Document(
            str(path)
        )

        parts = []

        # ----------------------------------------------------
        # Paragraphs
        # ----------------------------------------------------

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # ----------------------------------------------------
        # Tables
        # ----------------------------------------------------

        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    cell_text = (
                        cell.text.strip()
                    )

                    if cell_text:
                        cells.append(
                            cell_text
                        )

                if cells:
                    parts.append(
                        " | ".join(cells)
                    )

        full_text = "\n".join(
            parts
        ).strip()

    except Exception as error:
        return {
            "status": "error",
            "error": (
                f"Could not read DOCX: {error}"
            ),
        }

    if not full_text:
        return {
            "status": "error",
            "error": (
                "No extractable text was found "
                "in the DOCX file."
            ),
        }

    return {
        "status": "success",
        "text": full_text,
    }


# ============================================================
# INGEST TXT
# ============================================================

def ingest_txt_file(file_path):
    """
    Extract TXT and send it through
    the existing RAG ingestion pipeline.
    """

    extraction = extract_text_from_txt(
        file_path
    )

    if extraction.get("status") != "success":
        return extraction

    path = Path(file_path)

    metadata = {
        "file_name": path.name,
        "file_extension": path.suffix.lower(),
        "file_path": str(path),
    }

    return ingest_text(
        text=extraction["text"],
        source="txt_file",
        metadata=metadata,
    )


# ============================================================
# INGEST PDF
# ============================================================

def ingest_pdf_file(file_path):
    """
    Extract PDF and send it through
    the existing RAG ingestion pipeline.
    """

    extraction = extract_text_from_pdf(
        file_path
    )

    if extraction.get("status") != "success":
        return extraction

    path = Path(file_path)

    metadata = {
        "file_name": path.name,
        "file_extension": path.suffix.lower(),
        "file_path": str(path),
    }

    return ingest_text(
        text=extraction["text"],
        source="pdf_file",
        metadata=metadata,
    )


# ============================================================
# INGEST DOCX
# ============================================================

def ingest_docx_file(file_path):
    """
    Extract DOCX and send it through
    the existing RAG ingestion pipeline.
    """

    extraction = extract_text_from_docx(
        file_path
    )

    if extraction.get("status") != "success":
        return extraction

    path = Path(file_path)

    metadata = {
        "file_name": path.name,
        "file_extension": path.suffix.lower(),
        "file_path": str(path),
    }

    return ingest_text(
        text=extraction["text"],
        source="docx_file",
        metadata=metadata,
    )


# ============================================================
# GENERAL FILE INGESTION
# ============================================================

def ingest_file(file_path):
    """
    Detect the file type and send it
    to the appropriate ingestion function.
    """

    path = Path(file_path)

    if not path.exists():
        return {
            "status": "error",
            "error": "File does not exist.",
        }

    if not path.is_file():
        return {
            "status": "error",
            "error": "Path is not a file.",
        }

    extension = path.suffix.lower()

    if extension == ".txt":
        return ingest_txt_file(
            file_path
        )

    if extension == ".pdf":
        return ingest_pdf_file(
            file_path
        )

    if extension == ".docx":
        return ingest_docx_file(
            file_path
        )

    return {
        "status": "error",
        "error": (
            "Unsupported file type: "
            f"{extension}"
        ),
    }